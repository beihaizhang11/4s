"""
Excel管理模块
负责根据模板生成账单Excel文件
"""
import os
import shutil
from typing import Dict, List
from openpyxl import load_workbook
from openpyxl.workbook import Workbook


class ExcelManager:
    """Excel管理器类"""
    
    def __init__(self):
        pass
    
    @staticmethod
    def parse_cell_reference(cell_ref: str) -> tuple:
        """解析单元格引用（如"A1"）为(row, column)
        返回格式: (sheet_name, row, column) 或 (None, row, column)
        """
        # 支持格式: "A1" 或 "Sheet1!A1"
        if '!' in cell_ref:
            sheet_name, cell = cell_ref.split('!', 1)
        else:
            sheet_name = None
            cell = cell_ref
        
        # 分离列字母和行号
        col_letters = ""
        row_num = ""
        for char in cell:
            if char.isalpha():
                col_letters += char.upper()
            elif char.isdigit():
                row_num += char
        
        if not col_letters or not row_num:
            raise ValueError(f"无效的单元格引用: {cell_ref}")
        
        # 将列字母转换为列号
        col_num = 0
        for char in col_letters:
            col_num = col_num * 26 + (ord(char) - ord('A') + 1)
        
        return sheet_name, int(row_num), col_num
    
    def create_bill_from_template(
        self,
        template_path: str,
        output_path: str,
        data: Dict,
        field_mappings: Dict[str, List[str]]
    ) -> bool:
        """根据模板创建账单
        
        Args:
            template_path: 模板Excel文件路径
            output_path: 输出文件路径
            data: 要填充的数据字典
            field_mappings: 字段映射，格式: {"field_name": ["A1", "B2"], ...}
        
        Returns:
            是否成功
        """
        try:
            # 检查模板文件是否存在
            if not os.path.exists(template_path):
                raise FileNotFoundError(f"模板文件不存在: {template_path}")
            
            # 复制模板文件到输出路径
            output_dir = os.path.dirname(output_path)
            if output_dir and not os.path.exists(output_dir):
                os.makedirs(output_dir)
            
            shutil.copy2(template_path, output_path)
            
            # 加载工作簿
            wb = load_workbook(output_path)
            
            # 获取默认工作表
            default_sheet = wb.active
            
            # 遍历字段映射，填充数据
            for field_name, cell_refs in field_mappings.items():
                # 获取字段值
                field_value = data.get(field_name, "")
                
                # 将值转换为字符串
                if field_value is None:
                    field_value = ""
                else:
                    field_value = str(field_value)
                
                # 写入每个映射的单元格
                for cell_ref in cell_refs:
                    try:
                        sheet_name, row, col = self.parse_cell_reference(cell_ref)
                        
                        # 选择工作表
                        if sheet_name:
                            if sheet_name in wb.sheetnames:
                                sheet = wb[sheet_name]
                            else:
                                print(f"警告: 工作表 '{sheet_name}' 不存在，跳过单元格 {cell_ref}")
                                continue
                        else:
                            sheet = default_sheet
                        
                        # 写入单元格
                        sheet.cell(row=row, column=col, value=field_value)
                        
                    except Exception as e:
                        print(f"写入单元格 {cell_ref} 失败: {e}")
                        continue
            
            # 保存工作簿
            wb.save(output_path)
            wb.close()
            
            print(f"账单已生成: {output_path}")
            return True
            
        except Exception as e:
            print(f"生成账单失败: {e}")
            return False
    
    @staticmethod
    def get_column_letter(col_num: int) -> str:
        """将列号转换为列字母（如1->A, 27->AA）"""
        result = ""
        while col_num > 0:
            col_num -= 1
            result = chr(col_num % 26 + ord('A')) + result
            col_num //= 26
        return result
    
    @staticmethod
    def validate_cell_reference(cell_ref: str) -> bool:
        """验证单元格引用是否有效"""
        try:
            ExcelManager.parse_cell_reference(cell_ref)
            return True
        except:
            return False
